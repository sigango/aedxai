"""Export existing AED-XAI artifacts into a supervisor-ready results package.

This script is intentionally read-only with respect to the source result files.
It standardizes the existing per-image and per-detection CSVs, reports both
requested-image and valid-explanation cohorts, and writes CSV, Markdown, JSON,
and Excel outputs under ``results/supervisor_report`` by default.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DISPLAY_METHODS = {
    "aedxai": "AED-XAI (Ours)",
    "gradcam": "GradCAM",
    "gcame": "G-CAME",
    "dclose": "D-CLOSE",
    "lime": "LIME",
}
QUALITY_METRICS = (
    "pg",
    "ebpg",
    "oa",
    "oa_normalized",
    "insertion_auc",
    "deletion_auc",
    "sparsity",
    "composite",
)
SUPPLEMENTARY_METRICS = ("evaluator_time_sec",)
SELECTOR_METHODS = ("gradcam", "gcame", "dclose", "lime")


@dataclass(frozen=True)
class ExportConfig:
    results_root: Path
    output_dir: Path


def _project_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _git_value(*args: str) -> str:
    completed = subprocess.run(
        ["git", *args],
        cwd=PROJECT_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    return completed.stdout.strip() if completed.returncode == 0 else "unavailable"


def _safe_bool(series: pd.Series) -> pd.Series:
    if pd.api.types.is_bool_dtype(series):
        return series.fillna(False)
    return series.astype(str).str.lower().isin({"true", "1", "yes"})


def _baseline_method(path: Path) -> str:
    stem = path.stem.removeprefix("baseline_").removesuffix("_results")
    for method_key in SELECTOR_METHODS:
        if stem == method_key or stem.endswith(f"_{method_key}"):
            return DISPLAY_METHODS[method_key]
    raise ValueError(f"Cannot infer baseline method from {path}")


def _standardize_headline_file(path: Path, method: str) -> pd.DataFrame:
    raw = pd.read_csv(path)
    required = {
        "detector",
        "image_path",
        "num_detections",
        "composite_score",
        "mean_pg",
        "mean_ebpg",
        "mean_oa",
        "mean_sparsity",
        "mean_insertion_auc",
        "mean_deletion_auc",
        "total_time",
        "has_error",
    }
    missing = sorted(required - set(raw.columns))
    if missing:
        raise ValueError(f"{path} is missing required columns: {missing}")

    standardized = pd.DataFrame(
        {
            "dataset": "COCO val2017",
            "detector": raw["detector"].astype(str),
            "method": method,
            "image_id": raw["image_path"].map(lambda value: Path(str(value)).name),
            "image_path": raw["image_path"].astype(str),
            "num_detections": pd.to_numeric(raw["num_detections"], errors="coerce").fillna(0).astype(int),
            "pg": pd.to_numeric(raw["mean_pg"], errors="coerce"),
            "ebpg": pd.to_numeric(raw["mean_ebpg"], errors="coerce"),
            "oa": pd.to_numeric(raw["mean_oa"], errors="coerce"),
            "insertion_auc": pd.to_numeric(raw["mean_insertion_auc"], errors="coerce"),
            "deletion_auc": pd.to_numeric(raw["mean_deletion_auc"], errors="coerce"),
            "sparsity": pd.to_numeric(raw["mean_sparsity"], errors="coerce"),
            "composite": pd.to_numeric(raw["composite_score"], errors="coerce"),
            "evaluator_time_sec": pd.to_numeric(raw["total_time"], errors="coerce"),
            "iterations": pd.to_numeric(raw.get("iterations", 1), errors="coerce"),
            "converged": _safe_bool(raw.get("converged", pd.Series(False, index=raw.index))),
            "has_error": _safe_bool(raw["has_error"]),
            "source_file": str(path.relative_to(PROJECT_ROOT)),
        }
    )
    standardized["oa_normalized"] = ((standardized["oa"] + 1.0) / 2.0).clip(0.0, 1.0)
    standardized["valid_explanation"] = (
        (standardized["num_detections"] > 0)
        & ~standardized["has_error"]
        & standardized["composite"].notna()
    )
    if "explain_time" in raw.columns:
        standardized["explanation_time_sec"] = pd.to_numeric(raw["explain_time"], errors="coerce")
    else:
        standardized["explanation_time_sec"] = np.nan
    if "escalations" in raw.columns:
        standardized["escalations"] = pd.to_numeric(raw["escalations"], errors="coerce")
    else:
        standardized["escalations"] = np.nan
    if "selected_methods" in raw.columns:
        standardized["selected_methods"] = raw["selected_methods"].astype(str)
    else:
        standardized["selected_methods"] = ""
    return standardized


def load_headline(results_root: Path) -> tuple[pd.DataFrame, list[Path]]:
    sources: list[tuple[Path, str]] = []
    aed_paths = [
        results_root / "aedxai" / "results.csv",
        results_root / "aedxai_fasterrcnn_resnet50_fpn_v2" / "results.csv",
    ]
    for path in aed_paths:
        if path.exists():
            sources.append((path, DISPLAY_METHODS["aedxai"]))
    for path in sorted((results_root / "baseline").glob("baseline_*_results.csv")):
        sources.append((path, _baseline_method(path)))
    if not sources:
        raise FileNotFoundError(f"No headline result CSVs found under {results_root}")
    frames = [_standardize_headline_file(path, method) for path, method in sources]
    data = pd.concat(frames, ignore_index=True)
    return data, [path for path, _ in sources]


def load_cross_domain(results_root: Path) -> tuple[pd.DataFrame, list[Path]]:
    path = results_root / "cross_domain.csv"
    if not path.exists():
        return pd.DataFrame(), []
    raw = pd.read_csv(path)
    required = {
        "domain",
        "detector",
        "method",
        "image_id",
        "num_detections",
        "pg",
        "ebpg",
        "oa",
        "sparsity",
        "composite",
        "computation_time",
        "has_error",
    }
    missing = sorted(required - set(raw.columns))
    if missing:
        raise ValueError(f"{path} is missing required columns: {missing}")
    data = pd.DataFrame(
        {
            "dataset": raw["domain"].astype(str),
            "detector": raw["detector"].astype(str),
            "method": raw["method"].astype(str),
            "method_key": raw.get("method_key", "").astype(str),
            "image_id": raw["image_id"].astype(str),
            "image_path": raw.get("image_path", raw["image_id"]).astype(str),
            "num_detections": pd.to_numeric(raw["num_detections"], errors="coerce").fillna(0).astype(int),
            "pg": pd.to_numeric(raw["pg"], errors="coerce"),
            "ebpg": pd.to_numeric(raw["ebpg"], errors="coerce"),
            "oa": pd.to_numeric(raw["oa"], errors="coerce"),
            "insertion_auc": np.nan,
            "deletion_auc": np.nan,
            "sparsity": pd.to_numeric(raw["sparsity"], errors="coerce"),
            "composite": pd.to_numeric(raw["composite"], errors="coerce"),
            "evaluator_time_sec": pd.to_numeric(raw["computation_time"], errors="coerce"),
            "iterations": pd.to_numeric(raw.get("iterations", 1), errors="coerce"),
            "converged": _safe_bool(raw.get("converged", pd.Series(False, index=raw.index))),
            "has_error": _safe_bool(raw["has_error"]),
            "scene_complexity": raw.get("scene_complexity", "unknown").astype(str),
            "error": raw.get("error", "").fillna("").astype(str),
            "source_file": str(path.relative_to(PROJECT_ROOT)),
        }
    )
    data["oa_normalized"] = ((data["oa"] + 1.0) / 2.0).clip(0.0, 1.0)
    data["valid_explanation"] = (
        (data["num_detections"] > 0)
        & ~data["has_error"]
        & data["composite"].notna()
    )
    return data, [path]


def load_detector_comparison(results_root: Path) -> tuple[pd.DataFrame, list[Path]]:
    path = results_root / "compare_detectors_per_image.csv"
    if not path.exists():
        return pd.DataFrame(), []
    return _standardize_headline_file(path, "AED-XAI detector comparison"), [path]


def _describe(values: pd.Series) -> dict[str, float | int]:
    numeric = pd.to_numeric(values, errors="coerce").dropna().astype(float)
    n = int(len(numeric))
    if n == 0:
        return {
            "n_metric": 0,
            "mean": np.nan,
            "std": np.nan,
            "sem": np.nan,
            "ci95_low": np.nan,
            "ci95_high": np.nan,
            "median": np.nan,
            "q25": np.nan,
            "q75": np.nan,
            "min": np.nan,
            "max": np.nan,
        }
    mean = float(numeric.mean())
    std = float(numeric.std(ddof=1)) if n > 1 else 0.0
    sem = std / math.sqrt(n) if n else np.nan
    margin = 1.96 * sem
    return {
        "n_metric": n,
        "mean": mean,
        "std": std,
        "sem": sem,
        "ci95_low": mean - margin,
        "ci95_high": mean + margin,
        "median": float(numeric.median()),
        "q25": float(numeric.quantile(0.25)),
        "q75": float(numeric.quantile(0.75)),
        "min": float(numeric.min()),
        "max": float(numeric.max()),
    }


def coverage_summary(data: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for keys, group in data.groupby(group_cols, sort=True, dropna=False):
        key_values = keys if isinstance(keys, tuple) else (keys,)
        valid = group["valid_explanation"].astype(bool)
        is_aed = group["method"].eq(DISPLAY_METHODS["aedxai"])
        row: dict[str, Any] = dict(zip(group_cols, key_values))
        row.update(
            {
                "n_requested": int(len(group)),
                "n_valid_explanation": int(valid.sum()),
                "n_empty_or_invalid": int((~valid).sum()),
                "n_errors": int(group["has_error"].sum()),
                "coverage_rate": float(valid.mean()) if len(group) else np.nan,
                "detections_mean_requested": float(group["num_detections"].mean()),
                "detections_total": int(group["num_detections"].sum()),
                "iterations_mean": float(group["iterations"].mean()),
                "convergence_rate_aed_only": (
                    float(group.loc[is_aed, "converged"].mean()) if is_aed.any() else np.nan
                ),
            }
        )
        rows.append(row)
    return pd.DataFrame(rows)


def metric_summary_long(
    data: pd.DataFrame,
    group_cols: list[str],
    metrics: Iterable[str],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    metrics = [metric for metric in metrics if metric in data.columns]
    for keys, group in data.groupby(group_cols, sort=True, dropna=False):
        key_values = keys if isinstance(keys, tuple) else (keys,)
        key_row = dict(zip(group_cols, key_values))
        valid = group["valid_explanation"].astype(bool)
        counts = {
            "n_requested": int(len(group)),
            "n_valid_explanation": int(valid.sum()),
            "coverage_rate": float(valid.mean()) if len(group) else np.nan,
        }
        for metric in metrics:
            source = pd.to_numeric(group[metric], errors="coerce")
            for cohort, values in (
                ("requested_zero_fill", source.where(valid, 0.0).fillna(0.0)),
                ("valid_only", source.loc[valid]),
            ):
                row = {**key_row, **counts, "cohort": cohort, "metric": metric}
                row.update(_describe(values))
                rows.append(row)
    return pd.DataFrame(rows)


def metric_summary_wide(
    coverage: pd.DataFrame,
    long_summary: pd.DataFrame,
    group_cols: list[str],
) -> pd.DataFrame:
    selected = long_summary[
        long_summary["cohort"].isin({"requested_zero_fill", "valid_only"})
    ].copy()
    values = selected.pivot_table(
        index=group_cols,
        columns=["metric", "cohort"],
        values=["mean", "std", "ci95_low", "ci95_high"],
        aggfunc="first",
    )
    values.columns = [
        f"{metric}_{stat}_{'requested' if cohort == 'requested_zero_fill' else 'valid'}"
        for stat, metric, cohort in values.columns
    ]
    values = values.reset_index()
    return coverage.merge(values, on=group_cols, how="left")


def paired_against_aed(
    data: pd.DataFrame,
    group_cols: list[str],
    metrics: Iterable[str],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    metrics = [metric for metric in metrics if metric in data.columns]
    for keys, group in data.groupby(group_cols, sort=True, dropna=False):
        key_values = keys if isinstance(keys, tuple) else (keys,)
        key_row = dict(zip(group_cols, key_values))
        aed = group[group["method"] == DISPLAY_METHODS["aedxai"]].copy()
        if aed.empty:
            continue
        for baseline_method in sorted(set(group["method"]) - {DISPLAY_METHODS["aedxai"]}):
            baseline = group[group["method"] == baseline_method].copy()
            paired = aed.merge(
                baseline,
                on="image_id",
                suffixes=("_aed", "_baseline"),
                how="inner",
                validate="one_to_one",
            )
            if paired.empty:
                continue
            both_valid = paired["valid_explanation_aed"] & paired["valid_explanation_baseline"]
            for metric in metrics:
                aed_values = pd.to_numeric(paired[f"{metric}_aed"], errors="coerce")
                baseline_values = pd.to_numeric(paired[f"{metric}_baseline"], errors="coerce")
                aed_values = aed_values.where(paired["valid_explanation_aed"], 0.0).fillna(0.0)
                baseline_values = baseline_values.where(
                    paired["valid_explanation_baseline"], 0.0
                ).fillna(0.0)
                difference = aed_values - baseline_values
                description = _describe(difference)
                tolerance = 1e-12
                row = {
                    **key_row,
                    "baseline_method": baseline_method,
                    "metric": metric,
                    "n_common_images": int(len(paired)),
                    "n_both_valid": int(both_valid.sum()),
                    "aed_mean_requested": float(aed_values.mean()),
                    "baseline_mean_requested": float(baseline_values.mean()),
                    "mean_difference_aed_minus_baseline": description["mean"],
                    "std_difference": description["std"],
                    "ci95_difference_low": description["ci95_low"],
                    "ci95_difference_high": description["ci95_high"],
                    "aed_wins": int((difference > tolerance).sum()),
                    "ties": int((difference.abs() <= tolerance).sum()),
                    "aed_losses": int((difference < -tolerance).sum()),
                }
                rows.append(row)
    return pd.DataFrame(rows)


def build_selector_tables(
    results_root: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, list[Path]]:
    per_detection_frames: list[pd.DataFrame] = []
    method_rows: list[dict[str, Any]] = []
    oracle_rows: list[dict[str, Any]] = []
    label_rows: list[dict[str, Any]] = []
    sources: list[Path] = []

    for progress_path in sorted(results_root.glob("xai_selector_training_progress_*.csv")):
        detector = progress_path.stem.removeprefix("xai_selector_training_progress_")
        progress = pd.read_csv(progress_path)
        sources.append(progress_path)
        composite_columns: list[str] = []
        for method in SELECTOR_METHODS:
            required = [
                f"pg_{method}",
                f"ebpg_{method}",
                f"oa_{method}",
                f"sparsity_{method}",
                f"time_{method}",
            ]
            if not all(column in progress.columns for column in required):
                continue
            progress[f"oa_normalized_{method}"] = (
                (pd.to_numeric(progress[f"oa_{method}"], errors="coerce") + 1.0) / 2.0
            ).clip(0.0, 1.0)
            progress[f"composite_{method}"] = (
                pd.to_numeric(progress[f"ebpg_{method}"], errors="coerce")
                + progress[f"oa_normalized_{method}"]
                + pd.to_numeric(progress[f"sparsity_{method}"], errors="coerce")
            ) / 3.0
            composite_columns.append(f"composite_{method}")
            for metric in ("pg", "ebpg", "oa", "oa_normalized", "sparsity", "composite", "time"):
                method_rows.append(
                    {
                        "detector": detector,
                        "population": "all_progress_rows",
                        "method": DISPLAY_METHODS[method],
                        "metric": "evaluator_time_sec" if metric == "time" else metric,
                        **_describe(progress[f"{metric}_{method}"]),
                    }
                )

        if len(composite_columns) == len(SELECTOR_METHODS):
            matrix = progress[composite_columns].to_numpy(dtype=float)
            best_indices = np.nanargmax(matrix, axis=1)
            sorted_matrix = np.sort(matrix, axis=1)
            best_scores = sorted_matrix[:, -1]
            second_scores = sorted_matrix[:, -2]
            margins = best_scores - second_scores
            progress["best_method_name_equal_weight"] = [
                DISPLAY_METHODS[SELECTOR_METHODS[index]] for index in best_indices
            ]
            progress["oracle_composite_equal_weight"] = best_scores
            progress["oracle_margin_equal_weight"] = margins
            progress["oracle_is_ambiguous_margin_lt_0_02"] = margins < 0.02
            fixed_means = np.nanmean(matrix, axis=0)
            best_fixed_index = int(np.nanargmax(fixed_means))
            oracle_rows.append(
                {
                    "detector": detector,
                    "population": "all_progress_rows",
                    "n_detections": int(len(progress)),
                    "best_fixed_method": DISPLAY_METHODS[SELECTOR_METHODS[best_fixed_index]],
                    "best_fixed_composite": float(fixed_means[best_fixed_index]),
                    "oracle_composite": float(np.nanmean(best_scores)),
                    "oracle_headroom": float(np.nanmean(best_scores) - fixed_means[best_fixed_index]),
                    "n_ambiguous_margin_lt_0_02": int((margins < 0.02).sum()),
                    "ambiguous_fraction": float((margins < 0.02).mean()),
                }
            )
            labels = pd.Series(progress["best_method_name_equal_weight"]).value_counts()
            for method in (DISPLAY_METHODS[key] for key in SELECTOR_METHODS):
                count = int(labels.get(method, 0))
                label_rows.append(
                    {
                        "detector": detector,
                        "population": "all_progress_rows",
                        "method": method,
                        "count": count,
                        "fraction": count / len(progress) if len(progress) else np.nan,
                    }
                )
        progress.insert(0, "detector", detector)
        progress.insert(1, "population", "all_progress_rows")
        progress.insert(2, "source_file", str(progress_path.relative_to(PROJECT_ROOT)))
        per_detection_frames.append(progress)

    for final_path in sorted(results_root.glob("xai_selector_training_data_*.csv")):
        detector = final_path.stem.removeprefix("xai_selector_training_data_")
        final = pd.read_csv(final_path)
        sources.append(final_path)
        composite_columns = [
            f"composite_{method}" for method in SELECTOR_METHODS if f"composite_{method}" in final
        ]
        for method in SELECTOR_METHODS:
            for metric in ("pg", "ebpg", "oa", "sparsity", "composite", "time"):
                column = f"{metric}_{method}"
                if column not in final:
                    continue
                if metric == "oa":
                    method_rows.append(
                        {
                            "detector": detector,
                            "population": "retained_finalized_rows",
                            "method": DISPLAY_METHODS[method],
                            "metric": "oa_normalized",
                            **_describe(((pd.to_numeric(final[column], errors="coerce") + 1.0) / 2.0).clip(0, 1)),
                        }
                    )
                method_rows.append(
                    {
                        "detector": detector,
                        "population": "retained_finalized_rows",
                        "method": DISPLAY_METHODS[method],
                        "metric": "evaluator_time_sec" if metric == "time" else metric,
                        **_describe(final[column]),
                    }
                )
        if len(composite_columns) == len(SELECTOR_METHODS):
            matrix = final[composite_columns].to_numpy(dtype=float)
            best_scores = np.nanmax(matrix, axis=1)
            fixed_means = np.nanmean(matrix, axis=0)
            best_fixed_index = int(np.nanargmax(fixed_means))
            oracle_rows.append(
                {
                    "detector": detector,
                    "population": "retained_finalized_rows",
                    "n_detections": int(len(final)),
                    "best_fixed_method": DISPLAY_METHODS[SELECTOR_METHODS[best_fixed_index]],
                    "best_fixed_composite": float(fixed_means[best_fixed_index]),
                    "oracle_composite": float(np.nanmean(best_scores)),
                    "oracle_headroom": float(np.nanmean(best_scores) - fixed_means[best_fixed_index]),
                    "n_ambiguous_margin_lt_0_02": int(
                        pd.to_numeric(final.get("oracle_margin", np.nan), errors="coerce").lt(0.02).sum()
                    ),
                    "ambiguous_fraction": float(
                        pd.to_numeric(final.get("oracle_margin", np.nan), errors="coerce").lt(0.02).mean()
                    ),
                }
            )
        if "best_method_name" in final:
            labels = final["best_method_name"].map(
                lambda value: DISPLAY_METHODS.get(str(value).lower(), str(value))
            ).value_counts()
            for method in (DISPLAY_METHODS[key] for key in SELECTOR_METHODS):
                count = int(labels.get(method, 0))
                label_rows.append(
                    {
                        "detector": detector,
                        "population": "retained_finalized_rows",
                        "method": method,
                        "count": count,
                        "fraction": count / len(final) if len(final) else np.nan,
                    }
                )

    per_detection = (
        pd.concat(per_detection_frames, ignore_index=True) if per_detection_frames else pd.DataFrame()
    )
    return (
        per_detection,
        pd.DataFrame(method_rows),
        pd.DataFrame(oracle_rows),
        pd.DataFrame(label_rows),
        sources,
    )


def build_manifest(paths: Iterable[Path]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for path in sorted(set(paths)):
        frame = pd.read_csv(path)
        stat = path.stat()
        rows.append(
            {
                "source_file": str(path.relative_to(PROJECT_ROOT)),
                "sha256": _sha256(path),
                "modified_local": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
                "size_bytes": int(stat.st_size),
                "rows": int(len(frame)),
                "columns": int(len(frame.columns)),
                "column_names": " | ".join(frame.columns.astype(str).tolist()),
            }
        )
    return pd.DataFrame(rows)


def metric_definitions() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "metric": "PG",
                "column": "pg",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "Fraction of detections whose peak-saliency pixel lies inside the detector box.",
                "reporting_note": "Detector box is pseudo-ground truth; not part of the composite.",
            },
            {
                "metric": "EBPG",
                "column": "ebpg",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "Fraction of total saliency energy inside the detector box.",
                "reporting_note": "Detector box is pseudo-ground truth; contributes 1/3 of composite.",
            },
            {
                "metric": "Insertion AUC",
                "column": "insertion_auc",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "Area under confidence recovery as salient pixels are inserted.",
                "reporting_note": "Available in headline files, absent from cross_domain.csv.",
            },
            {
                "metric": "Deletion AUC",
                "column": "deletion_auc",
                "range": "[0, 1]",
                "direction": "lower",
                "definition": "Area under confidence as salient pixels are deleted.",
                "reporting_note": "Available in headline files, absent from cross_domain.csv.",
            },
            {
                "metric": "OA",
                "column": "oa",
                "range": "[-1, 1]",
                "direction": "higher",
                "definition": "Insertion AUC minus deletion AUC.",
                "reporting_note": "Mapped to (OA+1)/2 before entering composite.",
            },
            {
                "metric": "Normalized OA",
                "column": "oa_normalized",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "Linear normalization (OA+1)/2.",
                "reporting_note": "Contributes 1/3 of composite.",
            },
            {
                "metric": "Sparsity",
                "column": "sparsity",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "Gini coefficient of the saliency map.",
                "reporting_note": "Measures concentration, not correctness; contributes 1/3 of composite.",
            },
            {
                "metric": "Composite",
                "column": "composite",
                "range": "[0, 1]",
                "direction": "higher",
                "definition": "(EBPG + normalized OA + sparsity) / 3.",
                "reporting_note": "Paper-facing primary score in current config.",
            },
            {
                "metric": "Evaluator time",
                "column": "evaluator_time_sec",
                "range": "[0, infinity)",
                "direction": "lower",
                "definition": "Sum of evaluator computation_time over detections.",
                "reporting_note": "Legacy files exclude explanation, detector, VLM, and most pipeline wall time.",
            },
        ]
    )


def _markdown_table(frame: pd.DataFrame, columns: list[str], digits: int = 4) -> str:
    selected = frame.loc[:, [column for column in columns if column in frame.columns]].copy()
    integer_columns = {
        "n_requested",
        "n_valid_explanation",
        "n_common_images",
        "n_both_valid",
        "n_detections",
        "aed_wins",
        "ties",
        "aed_losses",
        "count",
    }
    for column in selected.select_dtypes(include=[np.number]).columns:
        if column in integer_columns:
            selected[column] = selected[column].map(
                lambda value: "" if pd.isna(value) else str(int(value))
            )
        else:
            selected[column] = selected[column].map(
                lambda value: "" if pd.isna(value) else f"{float(value):.{digits}f}"
            )
    headers = selected.columns.tolist()
    lines = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(["---"] * len(headers)) + "|",
    ]
    for row in selected.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(lines)


def write_report(
    path: Path,
    headline_wide: pd.DataFrame,
    headline_paired: pd.DataFrame,
    cross_wide: pd.DataFrame,
    cross_paired: pd.DataFrame,
    detector_comparison_wide: pd.DataFrame,
    selector_oracle: pd.DataFrame,
    manifest: pd.DataFrame,
    workbook_name: str,
) -> None:
    headline_columns = [
        "detector",
        "method",
        "n_requested",
        "n_valid_explanation",
        "coverage_rate",
        "pg_mean_requested",
        "ebpg_mean_requested",
        "oa_mean_requested",
        "insertion_auc_mean_requested",
        "deletion_auc_mean_requested",
        "sparsity_mean_requested",
        "composite_mean_requested",
    ]
    headline_table = _markdown_table(headline_wide, headline_columns)

    headline_composite = headline_paired[headline_paired["metric"] == "composite"].copy()
    paired_columns = [
        "detector",
        "baseline_method",
        "n_common_images",
        "mean_difference_aed_minus_baseline",
        "ci95_difference_low",
        "ci95_difference_high",
        "aed_wins",
        "ties",
        "aed_losses",
    ]
    paired_table = _markdown_table(headline_composite, paired_columns)

    cross_columns = [
        "dataset",
        "method",
        "n_requested",
        "n_valid_explanation",
        "coverage_rate",
        "pg_mean_valid",
        "ebpg_mean_valid",
        "oa_mean_valid",
        "sparsity_mean_valid",
        "composite_mean_valid",
        "composite_mean_requested",
    ]
    cross_table = _markdown_table(cross_wide, cross_columns)
    cross_composite = cross_paired[
        (cross_paired["metric"] == "composite")
        & (cross_paired["baseline_method"] == DISPLAY_METHODS["gcame"])
    ].copy()
    cross_paired_table = _markdown_table(
        cross_composite,
        [
            "dataset",
            "baseline_method",
            "n_common_images",
            "mean_difference_aed_minus_baseline",
            "ci95_difference_low",
            "ci95_difference_high",
            "aed_wins",
            "ties",
            "aed_losses",
        ],
    )
    oracle_table = _markdown_table(
        selector_oracle,
        [
            "detector",
            "population",
            "n_detections",
            "best_fixed_method",
            "best_fixed_composite",
            "oracle_composite",
            "oracle_headroom",
            "ambiguous_fraction",
        ],
    )
    detector_comparison_table = _markdown_table(
        detector_comparison_wide,
        [
            "detector",
            "n_requested",
            "n_valid_explanation",
            "coverage_rate",
            "pg_mean_valid",
            "ebpg_mean_valid",
            "oa_mean_valid",
            "insertion_auc_mean_valid",
            "deletion_auc_mean_valid",
            "sparsity_mean_valid",
            "composite_mean_valid",
        ],
    )

    earliest = manifest["modified_local"].min() if len(manifest) else "unknown"
    latest = manifest["modified_local"].max() if len(manifest) else "unknown"
    text = f"""# AED-XAI detailed evaluation results for supervisor

Generated: {datetime.now().astimezone().isoformat()}

Git commit: `{_git_value("rev-parse", "HEAD")}` (working tree has uncommitted changes: `{bool(_git_value("status", "--porcelain"))}`)

## Critical reporting caveat

- These tables describe the result CSVs currently present in `results/`; they are not a rerun.
- Source artifact modification dates span `{earliest}` to `{latest}`. The headline AED-XAI/baseline files predate the July selector-objective, cascade, timing, and reporting changes.
- The archived headline files contain no `escalations`, `selected_methods`, or `explain_time` columns. Do not present them as evidence for the current evaluate-then-escalate implementation.
- Faster R-CNN has no LIME headline artifact, so its method comparison is incomplete.
- `ablation_selector_size.csv` predates the July equal-weight oracle rebuild; its stored `oracle_score` scale is stale and should not be quoted as the current paper composite.
- `evaluator_time_sec` is legacy evaluator time dominated by insertion/deletion forwards. It is not explanation-generation or end-to-end latency.
- Detector boxes are pseudo-ground truth for PG/EBPG; no human explanation annotations or qualitative saliency outputs are included.
- Two cohorts are exported: `requested_zero_fill` counts an image with no valid explanation as zero; `valid_only` excludes it. Always state which denominator is reported.

## Headline COCO val2017 results

The table below uses all 1,000 requested images per available detector/method and zero-fills images without a valid explanation. Values are per-image means; detailed standard deviations, medians, quartiles, ranges, SEMs, and 95% normal-approximation CIs are in `headline_metrics_long.csv`.

{headline_table}

## Paired headline comparison against AED-XAI

Differences are `AED-XAI - baseline` on matched image IDs, using the requested-image zero-fill convention. A negative composite difference favors the baseline.

{paired_table}

## Cross-domain results

`composite_mean_valid` reproduces the historical summary convention that excludes no-detection rows. `composite_mean_requested` counts those rows as zero and is the coverage-aware value.

{cross_table}

## Cross-domain paired composite: AED-XAI versus G-CAME

{cross_paired_table}

## Separate detector-comparison artifact

`compare_detectors_per_image.csv` is a separate AED-XAI run. Its historical summary excluded no-detection rows, so the table below reports `valid_only`. It agrees with the conditional headline metrics but has different evaluator-time measurements.

{detector_comparison_table}

## Selector-training oracle diagnostic

These values are per detection and use the current equal-weight evaluator composite reconstructed from selector progress files. They are diagnostic oracle bounds, not headline test results.

{oracle_table}

## Safe supervisor statement

“In the archived 1,000-image COCO runs, fixed G-CAME has the highest mean composite for both YOLOX-S and Faster R-CNN. AED-XAI is lower on the same image IDs. The cross-domain archive shows the same ordering in all six YOLOX subsets. These are historical artifacts generated before the current selector/cascade changes, so a new provenance-controlled rerun is required before claiming performance for the current implementation.”

## Package contents

- `{workbook_name}`: all tables and normalized raw rows in one Excel workbook.
- `headline_per_image_raw.csv`: normalized 9,000-row headline image table.
- `headline_metrics_summary.csv`: compact detector/method summary.
- `headline_metrics_long.csv`: every metric with both cohorts and full descriptive statistics.
- `headline_pairwise_vs_aedxai.csv`: paired per-metric comparisons.
- `cross_domain_per_image_raw.csv`: normalized 6,000-row cross-domain table.
- `cross_domain_metrics_summary.csv`, `cross_domain_metrics_long.csv`, and `cross_domain_pairwise_vs_aedxai.csv`.
- `detector_comparison_per_image_raw.csv`, `detector_comparison_metrics_summary.csv`, and `detector_comparison_metrics_long.csv`.
- `selector_per_detection_raw.csv`: 4,000 selector-progress detections with reconstructed equal-weight composites and oracle labels.
- `selector_method_metrics.csv`, `selector_oracle_summary.csv`, and `selector_label_distribution.csv`.
- `metric_definitions.csv`: definitions, directions, and reporting caveats.
- `source_manifest.csv`: hashes, modification times, schemas, and row counts for every input.
- `export_metadata.json`: export provenance and row counts.
"""
    path.write_text(text, encoding="utf-8")


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column_cells in worksheet.iter_cols(
            min_row=1,
            max_row=min(worksheet.max_row, 250),
            min_col=1,
            max_col=worksheet.max_column,
        ):
            width = min(
                60,
                max(10, max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells) + 2),
            )
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def export(config: ExportConfig) -> dict[str, Any]:
    config.output_dir.mkdir(parents=True, exist_ok=True)

    headline, headline_sources = load_headline(config.results_root)
    cross, cross_sources = load_cross_domain(config.results_root)
    detector_comparison, detector_comparison_sources = load_detector_comparison(config.results_root)
    (
        selector_raw,
        selector_methods,
        selector_oracle,
        selector_labels,
        selector_sources,
    ) = build_selector_tables(config.results_root)

    headline_groups = ["dataset", "detector", "method"]
    headline_coverage = coverage_summary(headline, headline_groups)
    headline_long = metric_summary_long(
        headline,
        headline_groups,
        (*QUALITY_METRICS, *SUPPLEMENTARY_METRICS),
    )
    headline_wide = metric_summary_wide(headline_coverage, headline_long, headline_groups)
    headline_paired = paired_against_aed(
        headline,
        ["dataset", "detector"],
        (*QUALITY_METRICS, *SUPPLEMENTARY_METRICS),
    )

    cross_groups = ["dataset", "detector", "method"]
    if cross.empty:
        cross_coverage = pd.DataFrame()
        cross_long = pd.DataFrame()
        cross_wide = pd.DataFrame()
        cross_paired = pd.DataFrame()
    else:
        cross_coverage = coverage_summary(cross, cross_groups)
        cross_long = metric_summary_long(
            cross,
            cross_groups,
            ("pg", "ebpg", "oa", "oa_normalized", "sparsity", "composite", "evaluator_time_sec"),
        )
        cross_wide = metric_summary_wide(cross_coverage, cross_long, cross_groups)
        cross_paired = paired_against_aed(
            cross,
            ["dataset", "detector"],
            ("pg", "ebpg", "oa", "oa_normalized", "sparsity", "composite", "evaluator_time_sec"),
        )

    if detector_comparison.empty:
        detector_comparison_long = pd.DataFrame()
        detector_comparison_wide = pd.DataFrame()
    else:
        detector_comparison_groups = ["dataset", "detector", "method"]
        detector_comparison_coverage = coverage_summary(
            detector_comparison, detector_comparison_groups
        )
        detector_comparison_long = metric_summary_long(
            detector_comparison,
            detector_comparison_groups,
            (*QUALITY_METRICS, *SUPPLEMENTARY_METRICS),
        )
        detector_comparison_wide = metric_summary_wide(
            detector_comparison_coverage,
            detector_comparison_long,
            detector_comparison_groups,
        )

    definitions = metric_definitions()
    ablation_path = config.results_root / "ablation_selector_size.csv"
    ablation = pd.read_csv(ablation_path) if ablation_path.exists() else pd.DataFrame()
    sources = (
        headline_sources
        + cross_sources
        + detector_comparison_sources
        + selector_sources
    )
    if ablation_path.exists():
        sources.append(ablation_path)
    manifest = build_manifest(sources)

    csv_outputs = {
        "headline_per_image_raw.csv": headline,
        "headline_metrics_summary.csv": headline_wide,
        "headline_metrics_long.csv": headline_long,
        "headline_pairwise_vs_aedxai.csv": headline_paired,
        "cross_domain_per_image_raw.csv": cross,
        "cross_domain_metrics_summary.csv": cross_wide,
        "cross_domain_metrics_long.csv": cross_long,
        "cross_domain_pairwise_vs_aedxai.csv": cross_paired,
        "detector_comparison_per_image_raw.csv": detector_comparison,
        "detector_comparison_metrics_summary.csv": detector_comparison_wide,
        "detector_comparison_metrics_long.csv": detector_comparison_long,
        "selector_per_detection_raw.csv": selector_raw,
        "selector_method_metrics.csv": selector_methods,
        "selector_oracle_summary.csv": selector_oracle,
        "selector_label_distribution.csv": selector_labels,
        "ablation_selector_size_raw.csv": ablation,
        "metric_definitions.csv": definitions,
        "source_manifest.csv": manifest,
    }
    for filename, frame in csv_outputs.items():
        frame.to_csv(config.output_dir / filename, index=False)

    workbook_path = config.output_dir / "AEDXAI_supervisor_results.xlsx"
    readme = pd.DataFrame(
        [
            ("Scope", "Existing archived result artifacts; no model rerun was performed."),
            ("Primary cohort", "requested_zero_fill; valid_only is also provided."),
            ("Primary metric", "Composite = (EBPG + normalized OA + sparsity) / 3."),
            ("Timing warning", "Legacy evaluator time, not explanation or end-to-end latency."),
            ("Provenance warning", "Headline CSVs predate current selector/cascade code."),
            ("Missing result", "Faster R-CNN LIME headline result is absent."),
            ("Markdown report", "SUPERVISOR_RESULTS_REPORT.md"),
        ],
        columns=["item", "value"],
    )
    sheets = {
        "README": readme,
        "Headline_Summary": headline_wide,
        "Headline_Stats_Long": headline_long,
        "Headline_PerImage": headline,
        "Headline_Paired": headline_paired,
        "CrossDomain_Summary": cross_wide,
        "CrossDomain_Stats": cross_long,
        "CrossDomain_PerImage": cross,
        "CrossDomain_Paired": cross_paired,
        "DetectorCompare_Summary": detector_comparison_wide,
        "DetectorCompare_Raw": detector_comparison,
        "Selector_Methods": selector_methods,
        "Selector_Oracle": selector_oracle,
        "Selector_Labels": selector_labels,
        "Selector_PerDetection": selector_raw,
        "Ablation_Raw": ablation,
        "Source_Manifest": manifest,
        "Metric_Definitions": definitions,
    }
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    _format_workbook(workbook_path)

    report_path = config.output_dir / "SUPERVISOR_RESULTS_REPORT.md"
    write_report(
        report_path,
        headline_wide,
        headline_paired,
        cross_wide,
        cross_paired,
        detector_comparison_wide,
        selector_oracle,
        manifest,
        workbook_path.name,
    )

    metadata = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "git_commit": _git_value("rev-parse", "HEAD"),
        "working_tree_dirty": bool(_git_value("status", "--porcelain")),
        "results_root": str(config.results_root),
        "output_dir": str(config.output_dir),
        "row_counts": {filename: int(len(frame)) for filename, frame in csv_outputs.items()},
        "workbook": workbook_path.name,
        "report": report_path.name,
    }
    with (config.output_dir / "export_metadata.json").open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, indent=2)
    return metadata


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-root", default="results")
    parser.add_argument("--output-dir", default="results/supervisor_report")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    config = ExportConfig(
        results_root=_project_path(args.results_root),
        output_dir=_project_path(args.output_dir),
    )
    metadata = export(config)
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()
